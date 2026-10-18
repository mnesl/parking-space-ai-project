"""
Bonus Deliverable: Automated PowerPoint Presentation Generator (`scripts/generate_presentation.py`).

Generates `docs/AeroPark_AI_Reservation_Chatbot_Presentation.pptx` covering:
  - Slide 1: Title & Project Overview
  - Slide 2: High-Level Architecture (Hybrid RAG + Multi-Agent + MCP + LangGraph)
  - Slide 3: Stage 1 — Hybrid RAG (Pinecone + SQLite) & 3-Wall Sensitive Data Guardrails
  - Slide 4: Stage 1 — Scientific RAG Evaluation (`Recall@3`, `Precision@3`, Latency)
  - Slide 5: Stage 2 — Human-in-the-Loop (HITL) Administrator Agent & FastAPI Security
  - Slide 6: Stage 3 — Secure Model Context Protocol (MCP) Server & Delimiter Injection Defense
  - Slide 7: Stage 4 — LangGraph State Machine Orchestration & Concurrency Load Test Results
  - Slide 8: Live Demo Walkthrough & Screenshot Placeholders
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from pptx import Presentation
from pptx.util import Inches, Pt

from src.config import PROJECT_ROOT


def _add_bullet_slide(
    prs: Presentation,
    title_text: str,
    bullets: List[str],
) -> None:
    slide_layout = prs.slide_layouts[1]  # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title_text

    tf = slide.placeholders[1].text_frame
    tf.clear()
    for idx, line in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.space_after = Pt(10)


def generate_presentation(
    output_path: Path = PROJECT_ROOT
    / "docs"
    / "AeroPark_AI_Reservation_Chatbot_Presentation.pptx",
) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title Slide
    title_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_layout)
    slide1.shapes.title.text = (
        "AeroPark Smart Garage — Intelligent Parking Reservation Chatbot"
    )
    slide1.placeholders[1].text_frame.text = (
        "Hybrid RAG (Pinecone + SQLite) | Guardrails | Human-in-the-Loop | MCP Server | LangGraph\n"
        "Author: Merve Neslihan Okcu"
    )

    # Slide 2: System Architecture
    _add_bullet_slide(
        prs,
        "1. High-Level System Architecture (4 Stages)",
        [
            "LLM & Embeddings Factory: Provider-agnostic design supporting Google Gemini & OpenAI.",
            "Stage 1 (Hybrid RAG & Guardrails): Pinecone Serverless (Static Docs) + SQLite (Live Spaces/Prices).",
            "Stage 1 Security: 3-Wall Defense-in-Depth (Input Filter + Metadata Filter is_public=True + PII Redaction).",
            "Stage 2 (Human-in-the-Loop): Agent 2 (AdminEscalationAgent) + FastAPI REST Service with X-Admin-Api-Key.",
            "Stage 3 (Model Context Protocol): FastMCP + FastAPI Server writing 'Name | Car Number | Period | Approval Time'.",
            "Stage 4 (LangGraph Orchestration): 3-Node StateGraph with MemorySaver Checkpointer & Concurrency Load Testing.",
        ],
    )

    # Slide 3: Stage 1 — Hybrid RAG & Guardrails
    _add_bullet_slide(
        prs,
        "2. Stage 1: Hybrid RAG, Slot Filling & Sensitive Data Guardrails",
        [
            "Why Hybrid RAG? Static rules/directions live in Pinecone; changing space counts (58 spaces) live in SQLite.",
            "Chunking Strategy: RecursiveCharacterTextSplitter (chunk_size=550, chunk_overlap=80) preserving Markdown headers.",
            "Interactive Slot Filling: Pydantic ReservationDraft validates name, surname, car_number, reservation_period.",
            "3 Security Walls Against Prompt Injection & PII Exposure:",
            "   1. Wall #1 (Input Guardrail): Blocks jailbreaks and confidential credential/PII probes.",
            "   2. Wall #2 (Retrieval Filter & Scrubber): Enforces filter={'is_public': True} + Regex/NLP entity redaction.",
            "   3. Wall #3 (Output Guardrail): Scrubs SSNs, credit cards, tokens, and VIP PII before display.",
        ],
    )

    # Slide 4: Stage 1 — Evaluation Metrics
    _add_bullet_slide(
        prs,
        "3. Stage 1: Scientific RAG Evaluation Results",
        [
            "Evaluated across Golden Benchmark Queries + Adversarial Prompt Injection Suite:",
            "Mean Recall@3: 90.00% (Retrieved expected ground-truth documents inside Top-3 chunks).",
            "Mean Precision@3: 93.33% (High signal-to-noise ratio in retrieved context).",
            "Mean Response Factual Accuracy: 100.00% (Verified ground-truth facts in answers).",
            "Sensitive Data Guardrail Block Rate: 100.00% (0% leakage of SSNs, Credit Cards, or Override PINs).",
            "Median Latency (p50): 1.51 ms | 95th Percentile Latency (p95): 4.05 ms (local benchmark).",
        ],
    )

    # Slide 5: Stage 2 — Human-in-the-Loop
    _add_bullet_slide(
        prs,
        "4. Stage 2: Human-in-the-Loop (HITL) Administrator Agent",
        [
            "Multi-Agent Privilege Separation: Agent 1 (Customer Chatbot) cannot approve its own bookings.",
            "Agent 1 -> Agent 2 Handoff: Passes strictly validated ReservationDraft Pydantic model (no raw chat injection).",
            "Agent 2 (AdminEscalationAgent): Checks live SQLite zone capacity and generates an Executive Briefing.",
            "Asynchronous Persistence: Saves PENDING tickets in SQLite (REQ-XXXXXXXX) so server never blocks.",
            "FastAPI REST Endpoints (/admin/reservations/pending & /decision) secured by X-Admin-Api-Key (HTTP 401 on missing key).",
        ],
    )

    # Slide 6: Stage 3 — Secure MCP Server
    _add_bullet_slide(
        prs,
        "5. Stage 3: Model Context Protocol (MCP) Server & Security",
        [
            "Implemented via Official Python FastMCP SDK + FastAPI HTTP Transport + LangChain @tool fallback.",
            "Exact File Output Format in data/confirmed_reservations.txt:",
            "   Merve Okcu | 34ABC123 | 2026-10-20 09:00 to 2026-10-20 18:00 | 2026-10-18 22:05:00",
            "4-Layer Security & Reliability Hardening:",
            "   1. Status Gate: Rejects any ticket where status != 'CONFIRMED'.",
            "   2. Token Auth: Requires valid X-MCP-Api-Key header.",
            "   3. Delimiter Injection Prevention: Pydantic @field_validator blocks '|' and '\\n' characters.",
            "   4. Concurrency Safety: Uses threading.Lock() + flush() for atomic writes.",
        ],
    )

    # Slide 7: Stage 4 — LangGraph & Load Testing
    _add_bullet_slide(
        prs,
        "6. Stage 4: LangGraph Orchestration & Load Testing",
        [
            "LangGraph StateGraph connecting 3 Nodes via Conditional Edges:",
            "   Node 1 (user_interaction) -> Node 2 (admin_approval) -> Node 3 (data_recording).",
            "Checkpointer (MemorySaver): Freezes thread state at AWAITING_ADMIN_APPROVAL and resumes on Admin decision.",
            "Concurrent Load Test Results (30 requests x 6 concurrent worker threads per component):",
            "   • Stage 1 Chatbot RAG: 100% Success | 400.3 req/s | p50: 9.36 ms",
            "   • Stage 2 Admin Confirmation: 100% Success | 38.6 req/s | p50: 50.77 ms",
            "   • Stage 3 MCP File Recording: 100% Success | 32.5 req/s | p50: 73.12 ms",
            "   • Stage 4 End-to-End LangGraph: 100% Success | 21.2 req/s | p50: 228.10 ms",
        ],
    )

    # Slide 8: Demo & DevOps Bonus Summary
    _add_bullet_slide(
        prs,
        "7. Live Demo Walkthrough & Bonus Engineering Artifacts",
        [
            "Interactive Demo (Streamlit Multi-Tab UI + CLI + FastAPI Swagger /docs):",
            "   [Insert Screenshot / Video: Tab 1 — Customer Chatting & Providing 4 Reservation Slots]",
            "   [Insert Screenshot / Video: Tab 2 — Human Administrator Reviewing Briefing & Clicking CONFIRM]",
            "   [Insert Screenshot / Video: Tab 3 — MCP Server Entry in confirmed_reservations.txt & Live SQL Table]",
            "Bonus Artifacts Included in Repository:",
            "   • 12+ Automated Pytest Unit & Integration Tests across all modules (100% Passing)",
            "   • GitHub Actions CI/CD (.github/workflows/ci.yml), Dockerfile, docker-compose.yml & Terraform (terraform/main.tf)",
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


if __name__ == "__main__":
    saved = generate_presentation()
    print(f"Presentation generated at: {saved}")

